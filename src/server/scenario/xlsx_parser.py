import openpyxl


_CY20_SKILL_LAYOUT = {
    16: ("会计", "法律"),
    17: ("人类学", "图书馆使用"),
    18: ("估价", "聆听"),
    19: ("考古学", "锁匠"),
    20: ("技艺", "机械维修"),
    21: ("技艺", "医学"),
    22: ("技艺", "博物学"),
    23: ("取悦", "导航"),
    24: ("攀爬", "神秘学"),
    25: ("计算机使用", "操作重型机械"),
    26: ("信用评级", "说服"),
    27: ("克苏鲁神话", "驾驶"),
    28: ("乔装", "精神分析"),
    29: ("闪避", "心理学"),
    30: ("汽车驾驶", "骑术"),
    31: ("电气维修", "科学"),
    32: ("电子学", "科学"),
    33: ("话术", "科学"),
    34: ("格斗", "妙手"),
    35: ("格斗", "侦查"),
    36: ("格斗", "潜行"),
    37: ("格斗", "生存"),
    38: ("射击", "游泳"),
    39: ("射击", "投掷"),
    40: ("射击", "追踪"),
    41: ("射击", "驯兽"),
    42: ("急救", "潜水"),
    43: ("历史", "爆破"),
    44: ("恐吓", "读唇"),
    45: ("跳跃", "催眠"),
    46: ("外语", "炮术"),
    47: ("外语", "学识"),
    48: ("外语", "自定义技能"),
    49: ("母语", ""),
}


def parse_xlsx_character(file_path: str) -> dict:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    formula_wb = openpyxl.load_workbook(file_path, data_only=False)
    try:
        ws = _first_sheet(wb, ["人物卡", "Sheet1"]) or wb.active
        simple_ws = _first_sheet(wb, ["简化卡"])
        formula_ws = formula_wb[ws.title]
        if _looks_like_cy20(ws):
            return _parse_cy20_sheet(ws, simple_ws, formula_ws)
        return _parse_key_value_sheet(ws)
    finally:
        wb.close()
        formula_wb.close()


def _first_sheet(wb, names: list[str]):
    for name in names:
        if name in wb.sheetnames:
            return wb[name]
    return None


def _looks_like_cy20(ws) -> bool:
    label = str(_cell(ws, "B3") or "")
    return bool(_cell(ws, "E3")) and ("姓名" in label or "Name" in label)


def _parse_key_value_sheet(ws) -> dict:
    data = {}
    for row in ws.iter_rows(min_row=1, max_row=50, values_only=False):
        if len(row) >= 2 and row[0].value is not None and row[1].value is not None:
            key = str(row[0].value).strip()
            val = row[1].value
            data[key] = val

    name = data.get("姓名", data.get("Name", "未知"))
    hp = _to_int(data.get("HP", data.get("生命", 10)))
    san = _to_int(data.get("SAN", data.get("理智", 50)))
    mp = _to_int(data.get("MP", data.get("魔法", 10)))
    luck = _to_int(data.get("LUCK", data.get("幸运", 50)))

    skills = {}
    skip_keys = {"HP", "SAN", "MP", "LUCK", "生命", "理智", "魔法", "幸运", "姓名", "Name"}
    for k, v in data.items():
        if k not in skip_keys and isinstance(v, (int, float)):
            skills[k] = int(v)

    return {
        "name": str(name),
        "hp": hp,
        "max_hp": hp,
        "san": san,
        "max_san": san,
        "mp": mp,
        "max_mp": mp,
        "luck": luck,
        "skills": skills,
        "raw": data,
    }


def _parse_cy20_sheet(ws, simple_ws=None, formula_ws=None) -> dict:
    hp = _to_int(_cell(ws, "E10"), default=10)
    max_hp = _to_int(_cell(ws, "G10"), default=hp)
    san = _to_int(_cell(ws, "N10"), default=50)
    max_san = _to_int(_cell(ws, "P10"), default=san)
    mp = _to_int(_cell(simple_ws, "I7"), default=10) if simple_ws else 10
    max_mp = _to_int(_cell(simple_ws, "J7"), default=mp) if simple_ws else mp
    luck = _to_int(_cell(simple_ws, "I9"), default=50) if simple_ws else 50

    raw = {
        "name": _text(_cell(ws, "E3")),
        "occupation": _text(_cell(ws, "E5")),
        "age": _cell(ws, "E6"),
        "sex": _text(_cell(ws, "M6")),
        "residence": _text(_cell(ws, "E7")),
        "hometown": _text(_cell(ws, "M7")),
    }
    return {
        "name": raw["name"] or "未知",
        "occupation": raw["occupation"],
        "age": _to_int(raw["age"], default=0),
        "sex": raw["sex"],
        "residence": raw["residence"],
        "hometown": raw["hometown"],
        "hp": hp,
        "max_hp": max_hp,
        "san": san,
        "max_san": max_san,
        "mp": mp,
        "max_mp": max_mp,
        "luck": luck,
        "skills": _parse_cy20_skills(ws, formula_ws),
        "raw": raw,
    }


def _parse_cy20_skills(ws, formula_ws=None) -> dict[str, int]:
    skills: dict[str, int] = {}
    for row in range(16, min(ws.max_row, 49) + 1):
        left_name, right_name = _CY20_SKILL_LAYOUT.get(row, ("", ""))
        _add_cy20_skill(
            skills,
            ws,
            formula_ws,
            row,
            primary_column=6,
            secondary_column=8,
            total_column=18,
            component_columns=(10, 12, 14, 16),
            fallback_name=left_name,
        )
        _add_cy20_skill(
            skills,
            ws,
            formula_ws,
            row,
            primary_column=28,
            secondary_column=30,
            total_column=40,
            component_columns=(32, 34, 36, 38),
            fallback_name=right_name,
        )
    return skills


def _add_cy20_skill(
    skills: dict[str, int],
    ws,
    formula_ws,
    row: int,
    primary_column: int,
    secondary_column: int,
    total_column: int,
    component_columns: tuple[int, ...],
    fallback_name: str,
) -> None:
    value = _cy20_skill_total(ws, formula_ws, row, total_column, component_columns)
    if value is None:
        return
    primary = _text(ws.cell(row=row, column=primary_column).value) or _text(
        formula_ws.cell(row=row, column=primary_column).value if formula_ws else None
    ) or fallback_name
    secondary = _text(ws.cell(row=row, column=secondary_column).value) or _text(
        formula_ws.cell(row=row, column=secondary_column).value if formula_ws else None
    )
    name = _skill_name(primary, secondary)
    if name:
        skills[name] = value


def _cy20_skill_total(
    ws,
    formula_ws,
    row: int,
    total_column: int,
    component_columns: tuple[int, ...],
) -> int | None:
    cached_total = _number_or_none(ws.cell(row=row, column=total_column).value)
    if cached_total is not None:
        return cached_total

    components = [
        _cy20_component_value(ws, formula_ws, row, column)
        for column in component_columns
    ]
    if all(value is None for value in components):
        return None
    return sum(value or 0 for value in components)


def _cy20_component_value(ws, formula_ws, row: int, column: int) -> int | None:
    cached_value = _number_or_none(ws.cell(row=row, column=column).value)
    if cached_value is not None:
        return cached_value
    if formula_ws is None:
        return None

    formula = _text(formula_ws.cell(row=row, column=column).value).upper()
    if formula == "=INT(DEX/2)":
        return _to_int(_cell(ws, "AA3"), default=0) // 2
    if formula == "=AG5":
        return _to_int(_cell(ws, "AG5"), default=0)
    if formula.startswith("=IF(ISBLANK(AU26),5,"):
        return 5
    return None


def _skill_name(primary: str, secondary: str) -> str:
    if primary and secondary:
        base = primary.rstrip("①②③④⑤⑥⑦⑧⑨⑩1234567890").rstrip(":：")
        if base != primary or primary.endswith((":", "：")):
            return f"{base}：{secondary}"
        return secondary
    return primary or secondary


def _cell(ws, address: str):
    if ws is None:
        return None
    return ws[address].value


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _to_int(val, default: int = 0) -> int:
    if isinstance(val, (int, float)):
        return int(val)
    try:
        return int(str(val).strip())
    except (ValueError, TypeError):
        return default


def _number_or_none(value) -> int | None:
    if isinstance(value, (int, float)):
        return int(value)
    return None
