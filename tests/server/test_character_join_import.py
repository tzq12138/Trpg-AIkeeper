import json

import openpyxl
from tests.server.conftest import setup_auth_test_data, create_room


def _make_cy20_like_xlsx(path: str, name: str = "阿尔伯特·格雷", occupation: str = "罪犯-独行罪犯"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "人物卡"
    ws["B3"] = "姓名"
    ws["E3"] = name
    ws["B5"] = "职业"
    ws["E5"] = occupation
    ws["B6"] = "年龄"
    ws["E6"] = 26
    ws["J6"] = "性别"
    ws["M6"] = "男"
    ws["B7"] = "住地"
    ws["E7"] = "阿卡姆"
    ws["J7"] = "故乡"
    ws["M7"] = "波士顿"
    ws["B10"] = "体力 Hit  Points"
    ws["E10"] = 10
    ws["G10"] = 10
    ws["K10"] = "理智 Sanity"
    ws["N10"] = 55
    ws["P10"] = 99
    ws["F15"] = "技能名称"
    ws["R15"] = "成功率 普通/困难/极限"
    ws["F20"] = "技艺①"
    ws["H20"] = "表演"
    ws["R20"] = 80
    ws["F33"] = "话术"
    ws["R33"] = 25
    ws["F42"] = "急救"
    ws["R42"] = 55
    simple = wb.create_sheet("简化卡")
    simple["H7"] = "MP"
    simple["I7"] = 11
    simple["J7"] = 11
    simple["H9"] = "LUCK"
    simple["I9"] = 60
    wb.save(path)


def _room_id(client, test_db) -> str:
    setup_auth_test_data(test_db)
    return create_room(client)["room_id"]


def test_preview_xlsx_returns_summary_without_creating_character(client, test_db, tmp_path):
    path = tmp_path / "albert.xlsx"
    _make_cy20_like_xlsx(str(path))

    with path.open("rb") as file:
        resp = client.post(
            "/api/player/character/preview-xlsx",
            files={"file": ("albert.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert resp.status_code == 200
    data = resp.json()
    assert data["name"] == "阿尔伯特·格雷"
    assert data["occupation"] == "罪犯-独行罪犯"
    assert data["hp"] == 10
    assert data["max_hp"] == 10
    assert data["skill_count"] == 3

    row = test_db.execute("SELECT COUNT(*) AS c FROM characters").fetchone()
    assert row["c"] == 0


def test_join_with_uploaded_character_creates_player_and_character_data(client, test_db, tmp_path):
    room_id = _room_id(client, test_db)
    path = tmp_path / "albert.xlsx"
    _make_cy20_like_xlsx(str(path))

    with path.open("rb") as file:
        resp = client.post(
            f"/api/player/rooms/{room_id}/join-with-character",
            data={"player_name": "田文"},
            files={"file": ("albert.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert resp.status_code == 200
    joined = resp.json()
    assert joined["player_name"] == "田文"
    assert joined["investigator_name"] == "阿尔伯特·格雷"
    assert joined["character"]["skills"]["急救"] == 55

    char_resp = client.get(
        "/api/player/character",
        headers={"X-Room-Token": joined["player_token"]},
    )
    assert char_resp.status_code == 200
    char = char_resp.json()
    assert char["player_name"] == "田文"
    assert char["investigator_name"] == "阿尔伯特·格雷"
    assert char["name"] == "阿尔伯特·格雷"


def test_presets_list_skips_bad_files_and_marks_room_occupancy(client, test_db, tmp_path):
    room_id = _room_id(client, test_db)
    client.app.state.character_preset_dir = str(tmp_path)
    _make_cy20_like_xlsx(str(tmp_path / "albert.xlsx"))
    (tmp_path / "broken.xlsx").write_text("not a workbook", encoding="utf-8")

    resp = client.get(f"/api/player/character/presets?room_id={room_id}")

    assert resp.status_code == 200
    presets = resp.json()["presets"]
    assert len(presets) == 1
    assert presets[0]["preset_id"] == "albert"
    assert presets[0]["name"] == "阿尔伯特·格雷"
    assert presets[0]["occupied"] is False


def test_join_with_preset_locks_that_preset_for_room(client, test_db, tmp_path):
    room_id = _room_id(client, test_db)
    client.app.state.character_preset_dir = str(tmp_path)
    _make_cy20_like_xlsx(str(tmp_path / "albert.xlsx"))

    resp = client.post(
        f"/api/player/rooms/{room_id}/join-with-character",
        data={"player_name": "田文", "preset_id": "albert"},
    )

    assert resp.status_code == 200
    data = resp.json()
    assert data["player_name"] == "田文"
    assert data["investigator_name"] == "阿尔伯特·格雷"

    presets_resp = client.get(f"/api/player/character/presets?room_id={room_id}")
    assert presets_resp.json()["presets"][0]["occupied"] is True

    duplicate = client.post(
        f"/api/player/rooms/{room_id}/join-with-character",
        data={"player_name": "另一个玩家", "preset_id": "albert"},
    )

    assert duplicate.status_code == 409


def test_join_with_scenario_template_creates_initial_inventory(client, test_db):
    room_id = _room_id(client, test_db)
    test_db.execute(
        "INSERT INTO character_templates "
        "(template_id, scenario_id, name, occupation, attributes, skills, backstory) "
        "VALUES ('yhdx-reporter', 'sc-test', '查尔斯·钱伯斯', '记者', %s, %s, %s)",
        (
            json.dumps({"con": 55, "pow": 50, "luck": 50}),
            json.dumps({"侦查": 65, "图书馆使用": 70}),
            json.dumps({
                "inventory": [
                    {"name": "旅行箱", "description": "装着前往阿卡姆的行李。"},
                    {"name": "记者证", "description": "证明记者身份的证件。"},
                ]
            }, ensure_ascii=False),
        ),
    )
    test_db.commit()

    response = client.post(
        f"/api/player/rooms/{room_id}/join-with-character",
        data={"player_name": "田文", "template_id": "yhdx-reporter"},
    )

    assert response.status_code == 200
    rows = test_db.execute(
        "SELECT name, description, quantity, source FROM inventory "
        "WHERE character_id = %s ORDER BY acquired_at, name",
        (response.json()["character_id"],),
    ).fetchall()
    assert [row["name"] for row in rows] == ["旅行箱", "记者证"]
    assert all(row["quantity"] == 1 and row["source"] == "scenario_template" for row in rows)


def test_scenario_templates_endpoint_returns_public_template_fields(client, test_db):
    _room_id(client, test_db)
    test_db.execute(
        "INSERT INTO character_templates "
        "(template_id, scenario_id, name, occupation, background, age, gender, attributes, skills, backstory) "
        "VALUES ('public-template', 'sc-test', '艾达', '记者', '公开背景', 29, '女', %s, %s, %s)",
        (
            json.dumps({"con": 55, "pow": 60, "luck": 45}),
            json.dumps({"侦查": 65}),
            json.dumps({"inventory": [{"name": "隐藏物品"}], "secret": "不得返回"}),
        ),
    )
    test_db.commit()

    response = client.get("/api/scenarios/sc-test/templates")

    assert response.status_code == 200
    assert response.json() == [{
        "template_id": "public-template",
        "name": "艾达",
        "occupation": "记者",
        "background": "公开背景",
        "age": 29,
        "gender": "女",
        "attributes": {"con": 55, "pow": 60, "luck": 45},
        "skills": {"侦查": 65},
    }]


def test_join_with_multiple_sources_rejected(client, test_db, tmp_path):
    """Passing both preset_id and file should be rejected (400)."""
    room_id = _room_id(client, test_db)
    path = tmp_path / "albert.xlsx"
    _make_cy20_like_xlsx(str(path))
    with path.open("rb") as f:
        resp = client.post(
            f"/api/player/rooms/{room_id}/join-with-character",
            data={"player_name": "Test", "preset_id": "albert"},
            files={"file": ("albert.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert resp.status_code == 400


def test_join_triggers_lobby_snapshot(client, test_db, tmp_path):
    """Joining a room should write a s2c_room_lobby_snapshot event."""
    room_id = _room_id(client, test_db)
    path = tmp_path / "albert.xlsx"
    _make_cy20_like_xlsx(str(path))
    with path.open("rb") as f:
        resp = client.post(
            f"/api/player/rooms/{room_id}/join-with-character",
            data={"player_name": "Test"},
            files={"file": ("albert.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert resp.status_code == 200
    events = test_db.execute(
        "SELECT * FROM events WHERE room_id = %s AND event_type = 's2c_room_lobby_snapshot'",
        (room_id,),
    ).fetchall()
    assert len(events) >= 1, "Expected lobby snapshot event after join"


def test_failed_import_does_not_create_placeholder_player(client, test_db, tmp_path):
    room_id = _room_id(client, test_db)
    path = tmp_path / "broken.xlsx"
    path.write_text("not a workbook", encoding="utf-8")

    with path.open("rb") as file:
        resp = client.post(
            f"/api/player/rooms/{room_id}/join-with-character",
            data={"player_name": "田文"},
            files={"file": ("broken.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert resp.status_code == 400
    row = test_db.execute("SELECT COUNT(*) AS c FROM characters").fetchone()
    assert row["c"] == 0


def test_copy_character_wrong_account_rejected(client, test_db, tmp_path):
    """Cannot copy a character belonging to a different account."""
    from tests.server.conftest import create_account, create_scenario
    create_account(test_db, "acc-owner", "owneruser", "host")
    create_account(test_db, "acc-thief", "thiefuser", "player")
    create_scenario(test_db, "sc-copy", "Copy Test")
    test_db.commit()

    # Create room and character owned by acc-owner
    owner_token = client.post("/api/auth/login", json={
        "username": "owneruser", "password": "test123",
    }).json()["token"]
    room = client.post(
        "/api/rooms",
        json={"scenario_id": "sc-copy"},
        headers={"Authorization": f"Bearer {owner_token}"},
    ).json()
    path = tmp_path / "albert.xlsx"
    _make_cy20_like_xlsx(str(path))
    with path.open("rb") as f:
        join_resp = client.post(
            f"/api/player/rooms/{room['room_id']}/join-with-character",
            data={"player_name": "Owner"},
            files={"file": ("albert.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"Authorization": f"Bearer {owner_token}"},
        )
    assert join_resp.status_code == 200
    src_char_id = join_resp.json()["character_id"]

    # acc-thief tries to copy owner's character
    thief_token = client.post("/api/auth/login", json={
        "username": "thiefuser", "password": "test123",
    }).json()["token"]
    resp = client.post(
        f"/api/player/rooms/{room['room_id']}/join-with-character",
        data={"player_name": "Thief", "copy_character_id": src_char_id},
        headers={"Authorization": f"Bearer {thief_token}"},
    )
    assert resp.status_code == 403
